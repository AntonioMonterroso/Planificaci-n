#!/usr/bin/env python3
"""Extrae el TEXTO de un cuestionario/checklist .xlsx a un borrador JSON para revisar (nunca respuestas ni datos de clientes).

Uso:  python3 -I scripts/extraer_formato.py archivo.xlsx [salida.json]

Seguridad: lectura solo de valores (no ejecuta fórmulas ni macros). Rechaza .xlsm, avisa si el zip trae vbaProject.bin,
vínculos externos u objetos incrustados. Ejecútalo sobre copias descargadas en una carpeta aislada, no dentro del repo.
El resultado es un BORRADOR: se revisa a mano y luego se convierte en js/formatos/<id>.js.
"""
import json, re, sys, unicodedata, zipfile
from pathlib import Path

HEADER_WORDS = ("aplica", "cumple", "si/no", "sí/no", "respuesta")


def slug(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:60] or "formato"


def inspect_zip(path):
    warnings = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        warnings.append("Contiene macros VBA (vbaProject.bin): NO se ejecutan; revisar el código como texto si hace falta")
    if any("externallinks" in n.lower() for n in names):
        warnings.append("Tiene vínculos externos a otros libros")
    if any("/embeddings/" in n.lower() for n in names):
        warnings.append("Tiene objetos incrustados")
    return warnings


def is_question(text):
    t = text.strip()
    return len(t) > 12 and (t.startswith("¿") or t.endswith("?"))


def extract(path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise SystemExit("Solo se aceptan archivos .xlsx (los .xlsm traen macros y no se procesan)")
    warnings = inspect_zip(path)
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, keep_vba=False)
    sheets = []
    for ws in wb.worksheets:
        rows = [[(c if isinstance(c, str) else None) for c in r] for r in ws.iter_rows(values_only=True)]
        header_col = None
        for r in rows:
            for j, c in enumerate(r):
                if c and any(w in c.lower() for w in HEADER_WORDS) and len(c) < 30:
                    header_col = j if header_col is None else header_col
        items, section = [], None
        for i, r in enumerate(rows, start=1):
            texts = [(j, c.strip()) for j, c in enumerate(r) if c and c.strip()]
            if not texts:
                continue
            # la sección es una etiqueta corta en la primera columna con contenido a su derecha
            first_j, first = texts[0]
            if len(texts) > 1 and len(first) < 40 and not is_question(first) and first_j == 0:
                section = first
            for j, t in texts:
                if is_question(t) or (header_col is not None and j < header_col and len(t) > 25 and not t.isupper()):
                    items.append({"fila": i, "seccion": section, "texto": t})
                    break
        # quita duplicados exactos conservando el orden
        seen, uniq = set(), []
        for it in items:
            if it["texto"] not in seen:
                seen.add(it["texto"]); uniq.append(it)
        if uniq:
            sheets.append({"hoja": ws.title, "items": uniq})
    title = next((c.strip() for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for c in r if isinstance(c, str) and c.strip()), path.stem)
    return {"id": slug(path.stem), "fuente": path.name, "titulo": title, "borrador": True,
            "advertencias": warnings, "hojas": sheets, "total": sum(len(s["items"]) for s in sheets)}


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    out = extract(sys.argv[1])
    text = json.dumps(out, ensure_ascii=False, indent=2)
    if len(sys.argv) > 2:
        Path(sys.argv[2]).write_text(text, encoding="utf-8")
        print(f"{out['total']} ítems en {len(out['hojas'])} hoja(s) → {sys.argv[2]}" + (f" · {len(out['advertencias'])} advertencia(s)" if out["advertencias"] else ""))
    else:
        print(text)
